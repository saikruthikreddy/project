"""
Main document processing orchestrator that coordinates all file processors.
"""
import os
import logging
from pathlib import Path
from typing import List, Tuple, Dict, Any, Optional, Union
from dataclasses import dataclass

from giani_pkb.utils.exceptions import ProcessingError, FileProcessingError
from giani_pkb.utils.config import config
from giani_pkb.preprocessing.chunking import chunk_document_adaptive, ChunkMetadata

# Import new processors
from giani_pkb.preprocessing.pdf_processor import PdfProcessor
from giani_pkb.preprocessing.image_processor import ImageProcessor
from giani_pkb.preprocessing.csv_processor import CSVProcessor
from giani_pkb.preprocessing.pptx_processor import PptxProcessor
from giani_pkb.preprocessing.docx_processor import DocxProcessor

logger = logging.getLogger(__name__)

@dataclass
class ProcessingResult:
    """Result of processing a single document."""
    file_path: str
    document_id: str
    project_id: str
    status: str
    parsed_block_count: int
    chunk_count: int
    chunks_preview: List[Tuple[str, str, str]]  # (text_preview, chunk_id, chunk_type)
    error_message: Optional[str] = None

class DocumentProcessor:
    """
    Main orchestrator for document processing that coordinates all file processors.

    Features:
    - Unified interface for all document types
    - Automatic file type detection and routing
    - Configurable chunking strategies
    - Comprehensive error handling
    - Detailed processing results
    """

    def __init__(self, api_keys: Optional[Dict[str, str]] = None):
        """
        Initialize the document processor with all sub-processors.

        Args:
            api_keys: Dictionary of API keys for various services
        """
        self.api_keys = api_keys or {}

        # Initialize processors
        self._initialize_processors()

        # Supported file extensions
        self.supported_extensions = {
            # Images
            '.jpg', '.jpeg', '.png', '.bmp', '.tiff', '.tif', '.gif',
            # Documents
            '.pdf', '.docx', '.doc', '.pptx', '.ppt',
            # Data
            '.csv', '.xlsx', '.xls'
        }

    def _initialize_processors(self):
        """Initialize all document processors."""
        try:
            # Get API keys
            gemini_key = self.api_keys.get('gemini') or os.getenv('GEMINI_API_KEY') or getattr(config, 'GEMINI_API_KEY', None)
            openai_key = self.api_keys.get('openai') or os.getenv('OPENAI_API_KEY') or getattr(config, 'OPENAI_API_KEY', None)

            # Initialize processors
            self.pdf_processor = PdfProcessor()
            self.image_processor = ImageProcessor()
            self.csv_processor = CSVProcessor(api_key=gemini_key)
            self.pptx_processor = PptxProcessor(image_processor=self.image_processor)
            self.docx_processor = DocxProcessor()

            logger.info("All document processors initialized successfully")

        except Exception as e:
            logger.error(f"Error initializing processors: {e}")
            raise ProcessingError(f"Failed to initialize document processors: {e}")

    def _get_file_processor(self, file_extension: str):
        """
        Get the appropriate processor for a file extension.

        Args:
            file_extension: File extension (with dot)

        Returns:
            Appropriate processor instance

        Raises:
            ProcessingError: If file type is not supported
        """
        extension = file_extension.lower()

        if extension in ['.jpg', '.jpeg', '.png', '.bmp', '.tiff', '.tif', '.gif']:
            return self.image_processor
        elif extension == '.pdf':
            return self.pdf_processor
        elif extension in ['.csv', '.xlsx', '.xls']:
            return self.csv_processor
        elif extension in ['.pptx', '.ppt']:
            return self.pptx_processor
        elif extension in ['.docx', '.doc']:
            return self.docx_processor
        else:
            raise ProcessingError(f"Unsupported file type: {extension}")

    def _determine_document_type(self, document_category_hint: Optional[str] = None) -> str:
        """
        Determine document type for chunking based on category hint.

        Args:
            document_category_hint: Optional category hint

        Returns:
            Document type string for chunking
        """
        if document_category_hint:
            # Try to map category to document group
            mapped_group = CATEGORY_TO_GROUP_MAPPING.get(document_category_hint)
            if mapped_group:
                return mapped_group.value

        # Default document type
        return "formal"

    def _create_image_block(self, image_text: str) -> List[Tuple[str, Dict[str, Any]]]:
        """
        Create structured block from image processing result.

        Args:
            image_text: Text extracted from image

        Returns:
            List of structured blocks
        """
        if not image_text or not image_text.strip():
            return []

        return [(
            image_text.strip(),
            {
                "block_type": "image_full_text",
                "source_type": "image",
                "page_number": 1,
                "file_type": "image"
            }
        )]
    

    def process_file_light(self, file_path: Union[str, Path], max_chars: int = 5000) -> str:
        """
        Lightweight parsing for AI classification suggestions.
        Extracts ONLY plain text (no OCR, BLIP, or LlamaIndex).
        Truncates to first max_chars.
        """
        file_path = Path(file_path)
        if not file_path.exists():
            raise FileProcessingError(f"File not found: {file_path}", filepath=str(file_path))

        ext = file_path.suffix.lower()

        try:
            if ext == ".pdf":
                # Use PyMuPDF raw text only
                import fitz
                doc = fitz.open(str(file_path))
                text = []
                for page in doc:
                    text.append(page.get_text())
                    if len("".join(text)) > max_chars:
                        break
                return "".join(text)[:max_chars]

            elif ext in [".docx"]:
                import docx
                document = docx.Document(str(file_path))
                text = []
                for para in document.paragraphs:
                    text.append(para.text)
                    if len(" ".join(text)) > max_chars:
                        break
                return " ".join(text)[:max_chars]

            elif ext in [".pptx"]:
                from pptx import Presentation
                prs = Presentation(str(file_path))
                text = []
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if shape.has_text_frame:
                            text.append(shape.text.strip())
                        if len(" ".join(text)) > max_chars:
                            break
                return " ".join(text)[:max_chars]

            elif ext in ['.csv', '.xlsx', '.xls']:
                import pandas as pd
                if ext == ".csv":
                    df = pd.read_csv(str(file_path), nrows=50)  # only first rows
                else:
                    df = pd.read_excel(str(file_path), nrows=50)
                return df.to_csv(index=False)[:max_chars]

            elif ext in ['.txt', '.md']:
                return open(file_path, "r", encoding="utf-8", errors="ignore").read(max_chars)

            else:
                return ""  # unsupported extension for light parse

        except Exception as e:
            logger.error(f"Light parse failed for {file_path}: {e}")
            return ""

    def process_single_file(self,
                          file_path: Union[str, Path],
                          document_id: str,
                          project_id: str,
                          document_category_hint: Optional[str] = None,
                          use_semantic_chunker: bool = False,
                          **chunker_kwargs) -> Tuple[Optional[List[Tuple[str, Dict[str, Any]]]], List[Tuple[str, ChunkMetadata]]]:
        """
        Process a single file: parse it into blocks and then chunk those blocks.

        Args:
            file_path: Path to the file
            document_id: ID of the document
            project_id: ID of the project
            document_category_hint: Hint for document type classification
            use_semantic_chunker: Whether to use semantic chunking
            **chunker_kwargs: Additional chunking parameters

        Returns:
            Tuple of (parsed_blocks, chunks_with_metadata)

        Raises:
            FileProcessingError: If file processing fails
        """
        file_path = Path(file_path)

        if not file_path.exists():
            raise FileProcessingError(f"File not found: {file_path}", filepath=str(file_path))

        try:
            logger.info(f"Processing file: {file_path} for doc_id: {document_id}")

            # Get file extension and processor
            file_ext = file_path.suffix.lower()
            processor = self._get_file_processor(file_ext)

            # Process file based on type
            parsed_blocks = None

            if file_ext in ['.jpg', '.jpeg', '.png', '.bmp', '.tiff', '.tif', '.gif']:
                # Handle image processing
                result = processor.process_file(file_path)
                parsed_blocks = self._create_image_block(result['combined_text'])
            else:
                # Handle other file types
                parsed_blocks = processor.process_file(file_path)

            if not parsed_blocks:
                logger.warning(f"No content blocks extracted from {file_path} for doc_id: {document_id}")
                return None, []

            # Determine document type for chunking
            doc_type = self._determine_document_type(document_category_hint)
            logger.info(f"DocID {document_id}: Using document type '{doc_type}' for adaptive chunking.")

            # Get OpenAI API key for semantic chunking
            openai_key = self.api_keys.get('openai') or os.getenv('OPENAI_API_KEY')

            # Generate chunks
            chunks_with_metadata = chunk_document_adaptive(
                parsed_blocks,
                document_id,
                project_id,
                document_type=doc_type,
                use_semantic_chunker=use_semantic_chunker,
                openai_api_key=openai_key,
                **chunker_kwargs
            )

            logger.info(f"DocID {document_id}: Generated {len(chunks_with_metadata)} chunks.")
            return parsed_blocks, chunks_with_metadata

        except Exception as e:
            logger.error(f"Error processing file {file_path} for doc_id {document_id}: {e}", exc_info=True)
            raise FileProcessingError(f"Error processing file {file_path}: {e}", filepath=str(file_path))

    def process_files(self, file_metadata_list: List[Dict[str, Any]]) -> List[ProcessingResult]:
        """
        Process multiple files with their metadata.

        Args:
            file_metadata_list: List of file metadata dictionaries

        Returns:
            List of processing results
        """
        results = []

        for item in file_metadata_list:
            try:
                # Extract required metadata
                file_path = item.get("file_path")
                doc_id = item.get("document_id")
                proj_id = item.get("project_id")

                if not all([file_path, doc_id, proj_id]):
                    error_msg = "Missing critical metadata (file_path, document_id, or project_id)"
                    logger.error(f"{error_msg} in item: {item}")
                    results.append(ProcessingResult(
                        file_path=file_path or "unknown",
                        document_id=doc_id or "unknown",
                        project_id=proj_id or "unknown",
                        status="Error - Missing critical metadata",
                        parsed_block_count=0,
                        chunk_count=0,
                        chunks_preview=[],
                        error_message=error_msg
                    ))
                    continue

                # Extract optional parameters
                doc_category_hint = item.get("document_category_hint")
                use_semantic = item.get("use_semantic_chunker", False)
                chunker_settings = item.get("chunker_kwargs", {})

                # Process file
                parsed_blocks, chunks = self.process_single_file(
                    file_path, doc_id, proj_id,
                    document_category_hint=doc_category_hint,
                    use_semantic_chunker=use_semantic,
                    **chunker_settings
                )

                # Create result
                num_parsed_blocks = len(parsed_blocks) if parsed_blocks else 0
                chunks_preview = [
                    (chunk_text[:100] + "...", chunk_meta.chunk_id, chunk_meta.chunk_type)
                    for chunk_text, chunk_meta in chunks[:2]
                ]

                results.append(ProcessingResult(
                    file_path=file_path,
                    document_id=doc_id,
                    project_id=proj_id,
                    status="Success",
                    parsed_block_count=num_parsed_blocks,
                    chunk_count=len(chunks),
                    chunks_preview=chunks_preview
                ))

            except Exception as e:
                logger.error(f"Error processing item {item}: {e}")
                results.append(ProcessingResult(
                    file_path=item.get("file_path", "unknown"),
                    document_id=item.get("document_id", "unknown"),
                    project_id=item.get("project_id", "unknown"),
                    status="Error",
                    parsed_block_count=0,
                    chunk_count=0,
                    chunks_preview=[],
                    error_message=str(e)
                ))

        return results

    def discover_files(self, target_folder: Union[str, Path]) -> List[str]:
        """
        Discover all supported files in a directory.

        Args:
            target_folder: Path to search for files

        Returns:
            List of file paths
        """
        target_folder = Path(target_folder)

        if not target_folder.exists():
            logger.warning(f"Target folder does not exist: {target_folder}")
            return []

        file_paths = []
        try:
            for file_path in target_folder.rglob("*"):
                if file_path.is_file() and file_path.suffix.lower() in self.supported_extensions:
                    file_paths.append(str(file_path))

            logger.info(f"Discovered {len(file_paths)} supported files in {target_folder}")

        except Exception as e:
            logger.error(f"Error discovering files in {target_folder}: {e}")

        return file_paths

    def main_processing_orchestrator(self, file_metadata_list: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        """
        Legacy method for backward compatibility.

        Args:
            file_metadata_list: List of file metadata dictionaries

        Returns:
            List of result dictionaries
        """
        results = self.process_files(file_metadata_list)

        # Convert to legacy format
        legacy_results = []
        for result in results:
            legacy_results.append({
                "path": result.file_path,
                "document_id": result.document_id,
                "parsed_block_count": result.parsed_block_count,
                "chunk_count": result.chunk_count,
                "chunks_preview": result.chunks_preview
            })

        return legacy_results