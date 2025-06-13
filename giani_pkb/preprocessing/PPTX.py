# PPTX.py
from pptx import Presentation
import os
import tempfile # Added
from io import BytesIO
from PIL import Image as PILImage


class PPTX:
    def __init__(self, image_processor=None):
        self.image_processor = image_processor


    def process_pptx(self, file_path):
        try:
            presentation = Presentation(file_path)
            md_content = f"# {os.path.basename(file_path)}\n\n"
        
            for i, slide in enumerate(presentation.slides):
                md_content += f"## Slide {i + 1}\n\n"
            
                if slide.shapes.title and slide.shapes.title.has_text_frame:
                    title = slide.shapes.title.text
                    md_content += f"### {title}\n\n"
            
                for shape in slide.shapes:
                    if shape.has_text_frame and shape.text.strip():
                        if shape == slide.shapes.title:
                            continue
                        md_content += f"{shape.text}\n\n"
                    elif shape.shape_type == 13:
                        if self.image_processor:
                            temp_image_filepath = None # Initialize here for the finally block
                            try:
                                pil_image = self._extract_image_from_shape(shape)
                                if pil_image:
                                    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp_file:
                                        temp_image_filepath = tmp_file.name
                                    pil_image.save(temp_image_filepath, format="PNG")

                                    ocr_text_content = self.image_processor.process_image(temp_image_filepath)
                                    if ocr_text_content: # Make sure this is the text string
                                        md_content += f"**Image Text Content:**\n\n{ocr_text_content}\n\n"

                            except Exception as img_err:
                                md_content += f"*Error processing image: {str(img_err)}*\n\n"
                            finally:
                                if temp_image_filepath and os.path.exists(temp_image_filepath):
                                    os.remove(temp_image_filepath)
            
                md_content += "---\n\n"
        
            return md_content
        except Exception as e:
            return f"# Error Processing {os.path.basename(file_path)}\n\nAn error occurred: {str(e)}"


    def _extract_image_from_shape(self, shape):
        try:
            image_part = shape.image
            image_bytes = image_part.blob
            image = PILImage.open(BytesIO(image_bytes))
            return image
        except Exception as e:
            print(f"Error extracting image: {e}")
            return None
