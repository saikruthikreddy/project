import unittest
import os
from docx import Document as DocxDocument # Renamed to avoid clash with class name
from giani_pkb.preprocessing.Docx import Docx as DocxPreprocessor # Renamed for clarity

# Define the path for the dummy docx file within the tests directory
DUMMY_DOCX_FILENAME = "dummy_test_doc.docx"
# Assuming this test file is in tests/ directory, so path will be tests/dummy_test_doc.docx
DUMMY_DOCX_FILEPATH = os.path.join(os.path.dirname(__file__), DUMMY_DOCX_FILENAME)

class TestDocxPreprocessor(unittest.TestCase):
    def setUp(self):
        """Create a dummy DOCX file for testing."""
        self.doc_content = "Hello world. This is a test document.\nThis is the second line."

        # Create a new document
        doc = DocxDocument()
        doc.add_paragraph(self.doc_content)
        try:
            doc.save(DUMMY_DOCX_FILEPATH)
            # print(f"DEBUG: Dummy DOCX file created at {DUMMY_DOCX_FILEPATH}")
        except Exception as e:
            # print(f"DEBUG: Error creating dummy DOCX file: {e}")
            raise # Re-raise to make test setup failure clear

        self.preprocessor = DocxPreprocessor()

    def test_docx_preprocessor_happy_path(self):
        """Test processing a simple DOCX file."""
        if not os.path.exists(DUMMY_DOCX_FILEPATH):
            # print(f"DEBUG: Dummy file does not exist at path: {DUMMY_DOCX_FILEPATH} in test_docx_preprocessor_happy_path")
            # This might happen if setUp failed silently or if there's a path issue.
            # Re-create it for this test to proceed, though it indicates a problem.
            self.setUp() # Try to recreate it.
            if not os.path.exists(DUMMY_DOCX_FILEPATH):
                 self.fail(f"Setup failed to create dummy file at {DUMMY_DOCX_FILEPATH}")


        extracted_text = self.preprocessor.process_docx(DUMMY_DOCX_FILEPATH)

        # The python-docx library might add an extra newline if the text doesn't end with one,
        # or handle paragraphs by adding newlines.
        # We expect the content to be present.
        # For precise matching, compare paragraph texts.
        # However, process_docx joins paragraphs with "\n\n"

        # Let's read the doc again to see how python-docx structures it
        # doc = DocxDocument(DUMMY_DOCX_FILEPATH)
        # expected_text_from_paras = "\n\n".join([p.text for p in doc.paragraphs])
        # self.assertEqual(extracted_text.strip(), expected_text_from_paras.strip())

        # For this test, let's check for containment of the core content.
        # The `process_docx` method joins paragraphs with "\n\n".
        # If the original self.doc_content has "\n", it means it was a single paragraph with a soft break.
        # If it was two `add_paragraph` calls, then "\n\n" would be the separator.
        # In this setup, self.doc_content is passed as a single string to a single add_paragraph call.

        # The current Docx.process_docx joins paragraphs by "\n\n"
        # Since we added one paragraph, there should be no "\n\n".
        self.assertEqual(extracted_text.strip(), self.doc_content.strip())


    def tearDown(self):
        """Delete the dummy DOCX file."""
        try:
            if os.path.exists(DUMMY_DOCX_FILEPATH):
                os.remove(DUMMY_DOCX_FILEPATH)
                # print(f"DEBUG: Dummy DOCX file deleted from {DUMMY_DOCX_FILEPATH}")
        except Exception as e:
            # print(f"DEBUG: Error deleting dummy DOCX file: {e}")
            # Don't raise here, as it might obscure a test failure
            pass

if __name__ == '__main__':
    unittest.main()
