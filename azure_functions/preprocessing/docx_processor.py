"""
Azure-compatible DOCX file processor for extracting structured content, tables, and images.
Enhanced with LlamaIndex integration, TableTransformer, advanced OCR, and Azure Blob Storage support.
"""
import os
import docx
import zipfile
import io
import json
import uuid
import mimetypes
import re
import cv2
import numpy as np
from typing import List, Dict, Any, Tuple, Optional
import logging
from pathlib import Path
from datetime import datetime
from PIL import Image, ImageEnhance, ImageFilter

import pytesseract
try:
    from transformers import AutoProcessor, TableTransformerForObjectDetection
    import torch
    TRANSFORMER_AVAILABLE = True
except ImportError:
    TRANSFORMER_AVAILABLE = False
    logging.warning("TableTransformer not available. Using fallback methods.")

try:
    from llama_index.core.node_parser import SimpleNodeParser
    from llama_index.core.schema import Document as LlamaDocument
    from llama_index.core import Document, VectorStoreIndex
    LLAMAINDEX_AVAILABLE = True
except ImportError:
    LLAMAINDEX_AVAILABLE = False
    logging.warning("LlamaIndex not available. Enhanced parsing features will be disabled.")

from services.blob_storage_service import blob_storage_service
from utils.exceptions import ParsingError, FileProcessingError

logger = logging.getLogger(__name__)

class DocxProcessor:
    """
    Azure-compatible processor for DOCX files that extracts structured content, tables, and images.
    Enhanced with LlamaIndex integration, TableTransformer, and advanced OCR capabilities.

    Features:
    - Azure Blob Storage integration
    - Text extraction with style-based block classification
    - Advanced table detection using TableTransformer
    - Enhanced OCR with preprocessing for tables and charts
    - Chart data extraction and analysis
    - Image extraction and processing
    - Structured output with comprehensive metadata
    - LlamaIndex integration for enhanced document parsing
    - Node creation and document processing
    """

    def __init__(self, api_key: Optional[str] = None, use_llamaindex: bool = True):
        """
        Initialize the Azure-compatible DOCX processor.

        Args:
            api_key: Optional API key for image processing (if needed)
            use_llamaindex: Whether to use LlamaIndex for enhanced parsing
        """
        self.api_key = api_key
        self.use_llamaindex = use_llamaindex and LLAMAINDEX_AVAILABLE
        self.image_processor = self._get_image_processor()
        self.supported_extensions = {'.docx'}
        
        # Style mapping for better block classification
        self.style_mapping = {
            'heading 1': 'heading_1', 'heading 2': 'heading_2', 'heading 3': 'heading_3',
            'heading 4': 'heading_4', 'heading 5': 'heading_5', 'heading 6': 'heading_6',
            'list paragraph': 'list_item', 'caption': 'caption', 'quote': 'quote',
            'title': 'title', 'subtitle': 'subtitle'
        }
        
        # Setup advanced processing capabilities
        self._setup_advanced_processors()
        
        # Initialize LlamaIndex parser if available
        if self.use_llamaindex:
            self.node_parser = SimpleNodeParser.from_defaults()

    def _setup_advanced_processors(self):
        """Setup advanced processing models and configurations."""
        # Setup TableTransformer models
        if TRANSFORMER_AVAILABLE:
            try:
                self.table_detection_model = TableTransformerForObjectDetection.from_pretrained(
                    "microsoft/table-transformer-detection"
                )
                self.table_structure_model = TableTransformerForObjectDetection.from_pretrained(
                    "microsoft/table-transformer-structure-recognition"
                )
                self.table_processor = AutoProcessor.from_pretrained(
                    "microsoft/table-transformer-detection", use_fast=True
                )
                self.table_transformer_available = True
                logger.info("TableTransformer models loaded successfully")
            except Exception as e:
                logger.warning(f"Could not load TableTransformer models: {e}")
                self.table_transformer_available = False
        else:
            self.table_transformer_available = False

        # OCR configurations for different content types
        self.ocr_configs = {
            'table': '--oem 3 --psm 6 -c tessedit_char_whitelist=0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz.,()%$-+/= ',
            'chart': '--oem 3 --psm 11 -c tessedit_char_whitelist=0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz.,()%$-+/=:',
            'general': '--oem 3 --psm 6',
            'numbers': '--oem 3 --psm 8 -c tessedit_char_whitelist=0123456789.,%$-+'
        }

    def _get_image_processor(self):
        """Get image processor if available."""
        try:
            from preprocessing.image_processor import ImageProcessor
            return ImageProcessor()
        except ImportError:
            logger.warning("Image processor not available. Image extraction will be skipped.")
            return None

    def _get_page_numbers_from_element(self, element, doc: docx.Document) -> List[int]:
        """Extract page numbers for an element using document structure analysis."""
        try:
            page_breaks_before = 0
            current_element = element
            
            while current_element.getprevious() is not None:
                prev_element = current_element.getprevious()
                
                if hasattr(prev_element, 'tag') and 'sectPr' in prev_element.tag:
                    page_breaks_before += 1
                elif hasattr(prev_element, 'xpath'):
                    page_breaks = prev_element.xpath('.//w:br[@w:type="page"]', 
                                                   namespaces={'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'})
                    page_breaks_before += len(page_breaks)
                
                current_element = prev_element
            
            estimated_page = max(1, page_breaks_before + 1)
            return [estimated_page]
        except Exception as e:
            logger.warning(f"Could not determine page number: {e}")
            return [1]

    def _get_paragraph_style_type(self, paragraph: docx.text.paragraph.Paragraph) -> str:
        """
        Determine block type based on paragraph style.

        Args:
            paragraph: DOCX paragraph object

        Returns:
            Block type string
        """
        style_name = paragraph.style.name.lower()

        # Check for list items by content
        if paragraph.text.strip().startswith(("* ", "- ", "• ", "1. ", "2. ")):
            return "list_item"

        # Check style mapping
        for style_key, block_type in self.style_mapping.items():
            if style_key in style_name:
                return block_type

        # Check for generic headings
        if 'heading' in style_name:
            return "heading"

        return "paragraph"

    def ensure_rgb_image(self, image: Image.Image) -> Image.Image:
        """Ensure image is in RGB format for processing."""
        if image.mode != 'RGB':
            if image.mode == 'RGBA':
                background = Image.new('RGB', image.size, (255, 255, 255))
                background.paste(image, mask=image.split()[-1])
                return background
            else:
                return image.convert('RGB')
        return image

    def preprocess_image_for_ocr(self, image: Image.Image, content_type: str = 'general') -> Image.Image:
        """Preprocess image for better OCR results based on content type."""
        image = self.ensure_rgb_image(image)
        opencv_image = cv2.cvtColor(np.array(image), cv2.COLOR_RGB2BGR)
        
        # Content-specific preprocessing
        if content_type == 'chart':
            opencv_image = cv2.convertScaleAbs(opencv_image, alpha=1.5, beta=10)
            opencv_image = cv2.bilateralFilter(opencv_image, 9, 80, 80)
        elif content_type == 'table':
            kernel = np.array([[-1,-1,-1], [-1,9,-1], [-1,-1,-1]])
            opencv_image = cv2.filter2D(opencv_image, -1, kernel)
        
        # Convert to grayscale and apply adaptive thresholding
        gray = cv2.cvtColor(opencv_image, cv2.COLOR_BGR2GRAY)
        binary = cv2.adaptiveThreshold(gray, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C, cv2.THRESH_BINARY, 11, 2)
        
        # Clean up with morphological operations
        kernel = np.ones((1, 1), np.uint8)
        cleaned = cv2.morphologyEx(binary, cv2.MORPH_CLOSE, kernel)
        
        return Image.fromarray(cv2.cvtColor(cleaned, cv2.COLOR_GRAY2RGB))

    def detect_tables_with_transformer(self, image: Image.Image) -> List[Dict]:
        """Detect tables using TableTransformer or fallback to OCR."""
        if not self.table_transformer_available:
            return self.detect_tables_fallback(image)
        
        try:
            image = self.ensure_rgb_image(image)
            
            # Resize image if too large
            max_size = 1024
            if max(image.size) > max_size:
                ratio = max_size / max(image.size)
                new_size = tuple(int(dim * ratio) for dim in image.size)
                image = image.resize(new_size, Image.Resampling.LANCZOS)
            
            # Process with TableTransformer
            encoding = self.table_processor(image, return_tensors="pt")
            
            with torch.no_grad():
                outputs = self.table_detection_model(**encoding)
            
            target_sizes = torch.tensor([image.size[::-1]])
            results = self.table_processor.post_process_object_detection(
                outputs, threshold=0.7, target_sizes=target_sizes
            )[0]
            
            # Extract table data
            tables = []
            for score, label, box in zip(results["scores"], results["labels"], results["boxes"]):
                if score > 0.7:
                    box = [round(i, 2) for i in box.tolist()]
                    x1, y1, x2, y2 = box
                    
                    table_crop = image.crop((x1, y1, x2, y2))
                    table_data = self.extract_table_structure(table_crop)
                    
                    tables.append({
                        "bbox": box,
                        "confidence": score.item(),
                        "table_data": table_data,
                        "image_crop": table_crop
                    })
            
            return tables
            
        except Exception as e:
            logger.warning(f"TableTransformer detection failed: {e}, falling back to OCR")
            return self.detect_tables_fallback(image)

    def detect_tables_fallback(self, image: Image.Image) -> List[Dict]:
        """Fallback table detection using OCR and text analysis."""
        try:
            processed_img = self.preprocess_image_for_ocr(image, 'table')
            
            ocr_data = pytesseract.image_to_data(
                processed_img, 
                config=self.ocr_configs['table'],
                output_type=pytesseract.Output.DICT
            )
            
            # Extract text elements with positions
            text_lines = []
            words = ocr_data['text']
            confidences = ocr_data['conf']
            tops = ocr_data['top']
            lefts = ocr_data['left']
            
            for i, (word, conf, top, left) in enumerate(zip(words, confidences, tops, lefts)):
                if conf > 30 and word.strip():
                    text_lines.append({'text': word.strip(), 'top': top, 'left': left, 'confidence': conf})
            
            # Group by rows based on Y position
            rows = {}
            for item in text_lines:
                row_key = item['top'] // 20 * 20  # Group by approximate row
                if row_key not in rows:
                    rows[row_key] = []
                rows[row_key].append(item)
            
            # Check if this looks like a table (multiple rows with multiple columns)
            if len(rows) >= 2:
                avg_cols_per_row = sum(len(row) for row in rows.values()) / len(rows)
                if avg_cols_per_row >= 2:
                    table_text = self.format_fallback_table(rows)
                    return [{
                        "bbox": [0, 0, image.width, image.height],
                        "confidence": 0.6,
                        "table_data": {
                            "structure_detected": True,
                            "raw_text": table_text,
                            "parsed_rows": self.parse_table_text(table_text),
                            "cell_count": len(text_lines)
                        },
                        "image_crop": image
                    }]
            
            return []
            
        except Exception as e:
            logger.error(f"Fallback table detection failed: {e}")
            return []

    def format_fallback_table(self, rows: Dict) -> str:
        """Format OCR-detected table rows into readable text."""
        formatted_rows = []
        for top in sorted(rows.keys()):
            row_items = sorted(rows[top], key=lambda x: x['left'])
            row_text = " | ".join(item['text'] for item in row_items)
            formatted_rows.append(row_text)
        return "\n".join(formatted_rows)

    def extract_table_structure(self, table_image: Image.Image) -> Dict:
        """Extract structured data from a table image."""
        try:
            table_image = self.ensure_rgb_image(table_image)
            processed_img = self.preprocess_image_for_ocr(table_image, 'table')
            
            structure_detected = False
            cell_count = 0
            
            # Try structure recognition with TableTransformer
            if self.table_transformer_available:
                try:
                    encoding = self.table_processor(processed_img, return_tensors="pt")
                    
                    with torch.no_grad():
                        outputs = self.table_structure_model(**encoding)
                    
                    target_sizes = torch.tensor([processed_img.size[::-1]])
                    results = self.table_processor.post_process_object_detection(
                        outputs, threshold=0.5, target_sizes=target_sizes
                    )[0]
                    
                    structure_detected = len(results["boxes"]) > 0
                    cell_count = len(results["boxes"])
                    
                except Exception as e:
                    logger.warning(f"Structure recognition failed: {e}")
            
            # Extract text regardless of structure detection
            table_text = pytesseract.image_to_string(processed_img, config=self.ocr_configs['table'])
            rows = self.parse_table_text(table_text)
            
            return {
                "structure_detected": structure_detected,
                "raw_text": table_text,
                "parsed_rows": rows,
                "cell_count": cell_count if cell_count > 0 else len(rows) * (len(rows[0]) if rows else 0)
            }
            
        except Exception as e:
            logger.error(f"Table structure extraction failed: {e}")
            try:
                table_text = pytesseract.image_to_string(self.ensure_rgb_image(table_image), config=self.ocr_configs['general'])
                return {"structure_detected": False, "raw_text": table_text, "parsed_rows": [], "cell_count": 0}
            except Exception as fallback_e:
                logger.error(f"Fallback OCR also failed: {fallback_e}")
                return {"structure_detected": False, "raw_text": "", "parsed_rows": [], "cell_count": 0}

    def parse_table_text(self, text: str) -> List[List[str]]:
        """Parse table text into structured rows and columns."""
        lines = [line.strip() for line in text.split('\n') if line.strip()]
        rows = []
        
        for line in lines:
            # Split on multiple spaces or tabs to identify columns
            cells = re.split(r' {2,}|\t+', line)
            if len(cells) > 1:
                rows.append([cell.strip() for cell in cells])
        
        return rows

    def extract_chart_data_advanced(self, image: Image.Image) -> Dict[str, Any]:
        """Extract advanced chart data including type detection and structured elements."""
        try:
            image = self.ensure_rgb_image(image)
            processed_img = self.preprocess_image_for_ocr(image, 'chart')
            
            # Get OCR data with position information
            ocr_data = pytesseract.image_to_data(
                processed_img, 
                config=self.ocr_configs['chart'],
                output_type=pytesseract.Output.DICT
            )
            
            # Parse chart elements
            chart_elements = self.parse_chart_elements(ocr_data)
            
            # Get general text
            general_text = pytesseract.image_to_string(processed_img, config=self.ocr_configs['general'])
            
            # Detect chart type
            chart_type = self.detect_chart_type(general_text)
            
            return {
                "chart_type": chart_type,
                "structured_data": chart_elements,
                "raw_text": general_text,
                "confidence_score": self.calculate_chart_confidence(chart_elements, general_text)
            }
            
        except Exception as e:
            logger.error(f"Chart data extraction failed: {e}")
            return {"chart_type": "unknown", "structured_data": [], "raw_text": "", "confidence_score": 0.0}

    def parse_chart_elements(self, ocr_data: Dict) -> List[Dict]:
        """Parse OCR data to identify chart elements like labels and values."""
        elements = []
        words = ocr_data['text']
        confidences = ocr_data['conf']
        boxes = list(zip(ocr_data['left'], ocr_data['top'], ocr_data['width'], ocr_data['height']))
        
        for word, conf, box in zip(words, confidences, boxes):
            if conf < 30 or not word.strip():
                continue
                
            word = word.strip()
            
            # Identify numeric values and percentages
            if re.match(r'^\d+(?:\.\d+)?%?$', word):
                elements.append({
                    "type": "value", 
                    "text": word, 
                    "confidence": conf, 
                    "bbox": box, 
                    "is_percentage": '%' in word
                })
            # Identify labels
            elif len(word) > 2 and not word.isdigit():
                elements.append({
                    "type": "label", 
                    "text": word, 
                    "confidence": conf, 
                    "bbox": box
                })
        
        return elements

    def detect_chart_type(self, text: str) -> str:
        """Detect chart type based on text content."""
        text_lower = text.lower()
        
        if any(keyword in text_lower for keyword in ['pie', 'slice', 'sector']):
            return 'pie_chart'
        elif any(keyword in text_lower for keyword in ['bar', 'column', 'histogram']):
            return 'bar_chart'
        elif any(keyword in text_lower for keyword in ['line', 'trend', 'time series']):
            return 'line_chart'
        elif any(keyword in text_lower for keyword in ['scatter', 'plot', 'correlation']):
            return 'scatter_plot'
        else:
            return 'unknown'

    def calculate_chart_confidence(self, elements: List[Dict], text: str) -> float:
        """Calculate confidence score for chart detection."""
        if not elements:
            return 0.0
        
        avg_confidence = sum(elem['confidence'] for elem in elements) / len(elements)
        
        # Boost confidence if we have both labels and values
        has_labels = any(elem['type'] == 'label' for elem in elements)
        has_values = any(elem['type'] == 'value' for elem in elements)
        
        if has_labels and has_values:
            avg_confidence *= 1.2
        
        return min(avg_confidence / 100.0, 1.0)

    def _table_to_markdown(self, table: docx.table.Table) -> str:
        """
        Convert a DOCX table to Markdown format using enhanced logic.

        Args:
            table: DOCX table object

        Returns:
            Markdown table string
        """
        if not table.rows:
            return ""

        try:
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if not rows:
                return ""
            
            header = rows[0]
            separator = ["---"] * len(header)
            body = rows[1:]
            
            markdown = "| " + " | ".join(header) + " |\n"
            markdown += "| " + " | ".join(separator) + " |\n"
            
            for row in body:
                # Ensure row has same length as header
                padded_row = row + [""] * (len(header) - len(row))
                markdown += "| " + " | ".join(padded_row[:len(header)]) + " |\n"
            
            return markdown.strip()

        except Exception as e:
            logger.error(f"Error converting table to markdown: {e}")
            return self._table_to_simple_text(table)

    def _table_to_simple_text(self, table: docx.table.Table) -> str:
        """Convert table to simple text format as fallback."""
        text_lines = []
        for row in table.rows:
            row_texts = []
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    row_texts.append(cell_text)
            if row_texts:
                text_lines.append(" | ".join(row_texts))

        return "\n".join(text_lines)

    def _extract_image_info_enhanced(self, doc: docx.Document) -> List[Dict[str, Any]]:
        """
        Extract enhanced image information from DOCX document.

        Args:
            doc: DOCX document object

        Returns:
            List of image info dictionaries
        """
        image_info = []
        try:
            rels = doc.part.rels
            count = 0
            for rel in rels.values():
                if "image" in rel.reltype:
                    count += 1
                    try:
                        img_blob = rel.target_part.blob
                        try:
                            img = Image.open(io.BytesIO(img_blob))
                            width, height = img.size
                            format_type = img.format or "Unknown"
                            mode = img.mode
                            img_description = f"[IMAGE {count}: {format_type} image, {width}x{height} pixels, {mode} mode]"
                        except Exception as e:
                            img_description = f"[IMAGE {count}: Unable to analyze image - {str(e)}]"
                        
                        image_info.append({
                            "index": count,
                            "description": img_description,
                            "size": len(img_blob),
                            "blob": img_blob if self.image_processor else None
                        })
                    except Exception as e:
                        image_info.append({
                            "index": count,
                            "description": f"[IMAGE {count}: Error processing image - {str(e)}]",
                            "size": 0,
                            "blob": None
                        })
        except Exception as e:
            logger.error(f"Error extracting image info: {e}")
        
        return image_info

    def _extract_images_from_docx(self, container: str, blob_name: str) -> List[Tuple[str, Dict[str, Any]]]:
        """
        Extract images from DOCX file stored in Azure Blob Storage and process them.

        Args:
            container: Azure Blob Storage container name
            blob_name: Blob name for the DOCX file

        Returns:
            List of (image_text, metadata) tuples
        """
        if not self.image_processor:
            return []

        image_blocks: List[Tuple[str, Dict[str, Any]]] = []

        try:
            # Get DOCX file from Azure Blob Storage
            docx_bytes = blob_storage_service.download_blob(container, blob_name)
            
            with zipfile.ZipFile(io.BytesIO(docx_bytes), 'r') as docx_zip:
                media_files = [item for item in docx_zip.infolist() if item.filename.startswith('word/media/')]

                for media_file in media_files:
                    try:
                        image_bytes = docx_zip.read(media_file.filename)
                        img = Image.open(io.BytesIO(image_bytes))
                        img = self.ensure_rgb_image(img)
                        
                        # Try table detection first
                        detected_tables = self.detect_tables_with_transformer(img)
                        
                        if detected_tables:
                            # Process as table
                            for table_data in detected_tables:
                                table_text = self.format_table_data(table_data['table_data'])
                                
                                metadata = {
                                    "page_numbers": [1],
                                    "block_type": "table",
                                    "source_type": "table",
                                    "image_filename": os.path.basename(media_file.filename),
                                    "file_type": "docx_image_table",
                                    "extraction_confidence": table_data['confidence'],
                                    "table_detection_method": "transformer" if self.table_transformer_available else "ocr_fallback"
                                }
                                image_blocks.append((table_text.strip(), metadata))
                        else:
                            # Process as chart or general image
                            chart_data = self.extract_chart_data_advanced(img)
                            
                            if chart_data['raw_text'].strip():
                                metadata = {
                                    "page_numbers": [1],
                                    "block_type": "chart" if chart_data['chart_type'] != 'unknown' else "image_text",
                                    "source_type": "image",
                                    "image_filename": os.path.basename(media_file.filename),
                                    "file_type": "docx_image",
                                    "chart_type": chart_data['chart_type'],
                                    "extraction_confidence": chart_data['confidence_score']
                                }
                                image_blocks.append((chart_data['raw_text'].strip(), metadata))

                    except Exception as e:
                        logger.error(f"Error processing image {media_file.filename}: {e}")

        except Exception as e:
            logger.error(f"Error extracting images from DOCX {blob_name}: {e}")

        return image_blocks

    def format_table_data(self, table_data: Dict) -> str:
        """Format detected table data into markdown."""
        if table_data.get('parsed_rows'):
            rows = table_data['parsed_rows']
            if not rows:
                return table_data.get('raw_text', '')
            
            headers = rows[0] if rows else []
            if not headers:
                return table_data.get('raw_text', '')
                
            # Create markdown table
            markdown = "| " + " | ".join(headers) + " |\n"
            markdown += "| " + " | ".join(["---"] * len(headers)) + " |\n"
            
            for row in rows[1:]:
                padded_row = row + [''] * (len(headers) - len(row))
                markdown += "| " + " | ".join(padded_row[:len(headers)]) + " |\n"
            
            return markdown
        else:
            return table_data.get('raw_text', '')

    def _parse_docx_to_blocks_enhanced(self, container: str, blob_name: str) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
        """
        Parse DOCX to blocks using enhanced logic with Azure Blob Storage.

        Args:
            container: Azure Blob Storage container name
            blob_name: Blob name for the DOCX file

        Returns:
            Tuple of (blocks, image_info)
        """
        doc = blob_storage_service.get_docx_document(container, blob_name)
        blocks = []

        # Extract image information
        image_info_list = self._extract_image_info_enhanced(doc)
        
        # Create element mappings for efficient processing
        table_elements = {tbl._element: tbl for tbl in doc.tables}
        para_elements = {p._element: p for p in doc.paragraphs}

        img_counter = 0
        element_order = 0

        # Process document elements in order
        for element in doc.element.body:
            element_order += 1
            page_numbers = self._get_page_numbers_from_element(element, doc)
            
            # Process tables
            if element in table_elements:
                table = table_elements[element]
                markdown_table = self._table_to_markdown(table)
                if markdown_table:
                    blocks.append({
                        "type": "table",
                        "text": markdown_table,
                        "element_order": element_order,
                        "page_numbers": page_numbers,
                        "num_rows": len(table.rows),
                        "num_cols": len(table.columns) if table.rows else 0
                    })
            
            # Process paragraphs
            elif element in para_elements:
                para = para_elements[element]
                text = para.text.strip()
                if text:
                    blocks.append({
                        "type": "text",
                        "text": text,
                        "element_order": element_order,
                        "page_numbers": page_numbers,
                        "style_name": para.style.name,
                        "block_type": self._get_paragraph_style_type(para)
                    })
            
            # Process image/drawing elements
            elif element.tag.endswith("drawing"):
                if img_counter < len(image_info_list):
                    img_info = image_info_list[img_counter]
                    
                    try:
                        if img_info.get("blob"):
                            img = Image.open(io.BytesIO(img_info["blob"]))
                            img = self.ensure_rgb_image(img)
                            
                            # Try table detection first
                            detected_tables = self.detect_tables_with_transformer(img)
                            
                            if detected_tables:
                                for j, table_data in enumerate(detected_tables):
                                    table_text = self.format_table_data(table_data['table_data'])
                                    blocks.append({
                                        "type": "image_table",
                                        "text": table_text,
                                        "element_order": element_order + j * 0.1,
                                        "page_numbers": page_numbers,
                                        "image_size": img_info["size"],
                                        "image_index": img_info["index"],
                                        "table_confidence": table_data['confidence'],
                                        "extraction_method": "table_transformer" if self.table_transformer_available else "ocr_fallback"
                                    })
                            else:
                                # Process as chart
                                chart_data = self.extract_chart_data_advanced(img)
                                blocks.append({
                                    "type": "image_chart",
                                    "text": chart_data['raw_text'] or img_info["description"],
                                    "element_order": element_order,
                                    "page_numbers": page_numbers,
                                    "image_size": img_info["size"],
                                    "image_index": img_info["index"],
                                    "chart_type": chart_data['chart_type'],
                                    "chart_confidence": chart_data['confidence_score'],
                                    "structured_data": chart_data['structured_data']
                                })
                    except Exception as e:
                        logger.error(f"Enhanced image processing failed: {e}")
                        blocks.append({
                            "type": "image_reference",
                            "text": img_info["description"],
                            "element_order": element_order,
                            "page_numbers": page_numbers,
                            "image_size": img_info["size"],
                            "image_index": img_info["index"]
                        })
                    
                    img_counter += 1

        # Add any leftover images
        while img_counter < len(image_info_list):
            element_order += 1
            img_info = image_info_list[img_counter]
            blocks.append({
                "type": "image_reference",
                "text": img_info["description"],
                "element_order": element_order,
                "page_numbers": [1],
                "image_size": img_info["size"],
                "image_index": img_info["index"]
            })
            img_counter += 1

        return blocks, image_info_list

    def _process_paragraph(self, paragraph: docx.text.paragraph.Paragraph, element_order: int, page_numbers: List[int]) -> Optional[Tuple[str, Dict[str, Any]]]:
        """
        Process a single paragraph and return structured data.

        Args:
            paragraph: DOCX paragraph object
            element_order: Order of element in document
            page_numbers: Page numbers for this element

        Returns:
            (text, metadata) tuple or None if paragraph is empty
        """
        para_text = paragraph.text.strip()
        if not para_text:
            return None

        metadata = {
            "page_numbers": page_numbers,
            "block_type": self._get_paragraph_style_type(paragraph),
            "source_type": "text",
            "style_name": paragraph.style.name,
            "doc_element_order": element_order,
            "file_type": "docx"
        }

        return (para_text, metadata)

    def _process_table(self, table: docx.table.Table, element_order: int, page_numbers: List[int]) -> Optional[Tuple[str, Dict[str, Any]]]:
        """
        Process a single table and return structured data.

        Args:
            table: DOCX table object
            element_order: Order of element in document
            page_numbers: Page numbers for this element

        Returns:
            (text, metadata) tuple or None if table is empty
        """
        if not table.rows:
            return None

        table_text = self._table_to_markdown(table)
        if not table_text.strip():
            return None

        metadata = {
            "page_numbers": page_numbers,
            "block_type": "table",
            "source_type": "table",
            "num_rows": len(table.rows),
            "num_cols": len(table.columns),
            "doc_element_order": element_order,
            "file_type": "docx"
        }

        return (table_text, metadata)

    def create_llamaindex_nodes(self, blocks: List[Dict[str, Any]]) -> List[Any]:
        """
        Create LlamaIndex nodes from parsed blocks.

        Args:
            blocks: List of parsed blocks

        Returns:
            List of LlamaIndex nodes
        """
        if not self.use_llamaindex:
            logger.warning("LlamaIndex not available or disabled")
            return []

        try:
            nodes = []
            for block in blocks:
                doc = LlamaDocument(text=block["text"])
                doc.metadata = {
                    "page_numbers": block.get("page_numbers", [1]),
                    "block_type": block.get("block_type", "text"),
                    "source_type": block.get("source_type", "text"),
                    "element_order": block.get("element_order", 0)
                }
                
                block_nodes = self.node_parser.get_nodes_from_documents([doc])
                for node in block_nodes:
                    node.metadata.update(doc.metadata)
                    node.metadata["page_numbers"] = block.get("page_numbers", [1])
                nodes.extend(block_nodes)
            return nodes
        except Exception as e:
            logger.error(f"Error creating LlamaIndex nodes: {e}")
            return []

    def _normalize_blocks_input(self, blocks_input: Any) -> List[Dict[str, Any]]:
        """
        Normalize different types of blocks input to standard format.

        Args:
            blocks_input: Could be list of dicts, ProcessingResult, or other formats

        Returns:
            List of standardized block dictionaries
        """
        normalized_blocks = []
        
        if not blocks_input:
            return normalized_blocks
            
        # Handle list input
        if isinstance(blocks_input, list):
            for item in blocks_input:
                if isinstance(item, dict) and 'text' in item:
                    if 'page_numbers' not in item:
                        item['page_numbers'] = [1]
                    normalized_blocks.append(item)
                elif isinstance(item, str):
                    normalized_blocks.append({"text": item, "type": "text", "page_numbers": [1]})
                elif hasattr(item, 'text'):
                    normalized_blocks.append({
                        "text": str(item.text),
                        "type": getattr(item, 'type', 'text'),
                        "page_numbers": getattr(item, 'page_numbers', [1])
                    })
        
        # Handle iterable objects
        elif hasattr(blocks_input, '__iter__') and not isinstance(blocks_input, (str, dict)):
            try:
                for item in blocks_input:
                    if isinstance(item, dict) and 'text' in item:
                        if 'page_numbers' not in item:
                            item['page_numbers'] = [1]
                        normalized_blocks.append(item)
                    elif hasattr(item, 'text'):
                        normalized_blocks.append({
                            "text": str(item.text),
                            "type": getattr(item, 'type', 'text'),
                            "page_numbers": getattr(item, 'page_numbers', [1])
                        })
            except Exception as e:
                logger.warning(f"Error normalizing blocks input: {e}")
        
        # Handle ProcessingResult or similar objects
        elif hasattr(blocks_input, 'chunks_preview'):
            try:
                for chunk_info in blocks_input.chunks_preview[:2]:
                    if len(chunk_info) >= 1:
                        normalized_blocks.append({
                            "text": str(chunk_info[0]),
                            "type": "prose",
                            "page_numbers": [1]
                        })
            except Exception as e:
                logger.warning(f"Error extracting from ProcessingResult: {e}")
        
        return normalized_blocks

    def generate_metadata_with_llamaindex(self, container: str, blob_name: str, blocks: Any, nodes: List[Any] = None) -> Dict[str, Any]:
        """
        Generate comprehensive metadata using LlamaIndex integration and Azure Blob Storage.

        Args:
            container: Azure Blob Storage container name
            blob_name: Blob name for the DOCX file
            blocks: Parsed blocks (can be various formats)
            nodes: Optional LlamaIndex nodes

        Returns:
            Comprehensive metadata dictionary
        """
        document_id = str(uuid.uuid4())
        now_iso = datetime.now().isoformat()

        # Normalize blocks input to standard format
        actual_blocks = self._normalize_blocks_input(blocks)

        # Create preview from nodes or blocks
        preview_entries = []

        # Generate preview from nodes first (preferred)
        if nodes and len(nodes) > 0:
            for i, node in enumerate(nodes[:2]):
                try:
                    if hasattr(node, 'text'):
                        snippet = str(node.text).replace('\n', ' ')[:100].strip()
                        node_id = getattr(node, 'node_id', str(uuid.uuid4()))
                        page_nums = node.metadata.get('page_numbers', [1]) if hasattr(node, 'metadata') else [1]
                        preview_entries.append({
                            "snippet": snippet,
                            "node_id": node_id,
                            "type": "prose",
                            "page_numbers": page_nums
                        })
                except Exception as e:
                    logger.warning(f"Error processing node {i}: {e}")

        # Fallback to blocks if no valid nodes
        if not preview_entries and actual_blocks:
            for i, block in enumerate(actual_blocks[:2]):
                try:
                    text_content = block.get("text", "")
                    if text_content:
                        snippet = str(text_content).replace('\n', ' ')[:100].strip()
                        preview_entries.append({
                            "snippet": snippet,
                            "node_id": str(uuid.uuid4()),
                            "type": block.get("type", "prose"),
                            "page_numbers": block.get("page_numbers", [1])
                        })
                except Exception as e:
                    logger.warning(f"Error processing block {i}: {e}")

        # Final fallback - create a basic preview
        if not preview_entries:
            preview_entries = [{
                "snippet": "Document content available - preview generation failed",
                "node_id": str(uuid.uuid4()),
                "type": "prose",
                "page_numbers": [1]
            }]

        preview_str = json.dumps(preview_entries, ensure_ascii=False, indent=2)

        # Generate full text
        full_text = ""
        if nodes:
            try:
                full_text = "\n\n".join(str(node.text) for node in nodes if hasattr(node, 'text'))
            except Exception as e:
                logger.warning(f"Error generating full text from nodes: {e}")

        if not full_text and actual_blocks:
            try:
                full_text = "\n\n".join(block["text"] for block in actual_blocks if block.get("text"))
            except Exception as e:
                logger.warning(f"Error generating full text from blocks: {e}")
                full_text = "Content available - text generation failed"

        # Get file information from Azure Blob Storage
        try:
            file_info = blob_storage_service.get_blob_info(container, blob_name)
            file_size = file_info.get("size_human", "Unknown")
            mime_type = file_info.get("content_type", "")
        except Exception as e:
            logger.warning(f"Error getting blob info: {e}")
            file_size = "Unknown"
            mime_type = ""

        metadata = {
            "document_id": document_id,
            "dateAddedToGiani": now_iso,
            "originalFilename": Path(blob_name).name,
            "storagePath": f"{container}/{blob_name}",
            "fileSize": file_size,
            "fileMimeType": mime_type or "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "userID": "user_001",
            "projectID": "project_001",
            "categoryFolder": "Client-Provided Material",
            "finalCategory": "2. Operational Report/Review Deck",
            "finalPurpose": f"\"{blob_name}\" likely contains operational insights or structured data, including tables and embedded images, meant for review or reporting.",
            "priority": "Medium",
            "geminiPrompt": (
                "You are an AI assistant helping classify documents for a management consulting project.\n\n"
                f"This is the document source \"<user_input>preview</user_input>\".\n\n"
                f"Based on the filename \"<user_input>{blob_name}</user_input>\" and this text preview:\n"
                f"\"<user_input>{preview_str}</user_input>\"\n\nPlease classify this document..."
            ),
            "textPreview": preview_str,
            "storedFilename": f"data/uploaded_documents/Client-Provided Material/{blob_name}",
            "finalizedAt": now_iso,
            "savedAt": now_iso,
            "fullText": full_text,
            "parsedBlocksCount": len(actual_blocks),
            "llamaIndexNodesCount": len(nodes) if nodes else 0,
            "enhancedParsing": self.use_llamaindex,
            "advancedFeatures": {
                "tableTransformerUsed": self.table_transformer_available,
                "enhancedOcrUsed": True,
                "chartDetectionUsed": True,
                "structuredTableExtraction": True,
                "azureBlobStorage": True
            }
        }

        return metadata

    def process_file(self, container: str, blob_name: str, use_enhanced_parsing: bool = None) -> List[Tuple[str, Dict[str, Any]]]:
        """
        Process a DOCX file from Azure Blob Storage and extract structured content.

        Args:
            container: Azure Blob Storage container name
            blob_name: Blob name for the DOCX file
            use_enhanced_parsing: Override for using enhanced parsing

        Returns:
            List of (text_block, metadata) tuples

        Raises:
            ParsingError: If file format is not supported
            FileProcessingError: If processing fails
        """
        # Validate file extension
        file_extension = Path(blob_name).suffix.lower()
        if file_extension not in self.supported_extensions:
            raise ParsingError(
                f"Unsupported file format: {file_extension}. Supported: {', '.join(self.supported_extensions)}",
                filename=blob_name
            )

        # Check if blob exists
        if not blob_storage_service.blob_exists(container, blob_name):
            raise FileProcessingError(f"File not found in Azure Blob Storage: {container}/{blob_name}", filepath=blob_name)

        # Determine parsing method
        enhanced_parsing = use_enhanced_parsing if use_enhanced_parsing is not None else True

        try:
            if enhanced_parsing:
                return self._process_file_enhanced(container, blob_name)
            else:
                return self._process_file_legacy(container, blob_name)
        except Exception as e:
            logger.error(f"Error processing DOCX file {blob_name}: {e}")
            raise FileProcessingError(f"Error processing DOCX file {blob_name}: {e}", filepath=blob_name)

    def _process_file_enhanced(self, container: str, blob_name: str) -> List[Tuple[str, Dict[str, Any]]]:
        """
        Process file using enhanced logic with advanced features.

        Args:
            container: Azure Blob Storage container name
            blob_name: Blob name for the DOCX file

        Returns:
            List of (text_block, metadata) tuples
        """
        processed_blocks: List[Tuple[str, Dict[str, Any]]] = []

        # Parse using enhanced logic
        blocks, image_info = self._parse_docx_to_blocks_enhanced(container, blob_name)

        # Convert blocks to the expected format
        for block in blocks:
            metadata = {
                "page_numbers": block.get("page_numbers", [1]),
                "block_type": block.get("block_type", block["type"]),
                "source_type": self._determine_source_type(block["type"]),
                "doc_element_order": block.get("element_order", 0),
                "file_type": "docx"
            }

            # Add type-specific metadata
            if block["type"] == "table":
                metadata.update({
                    "num_rows": block.get("num_rows"),
                    "num_cols": block.get("num_cols")
                })
            elif block["type"] == "image_table":
                metadata.update({
                    "source_type": "table",
                    "block_type": "table",
                    "extraction_method": block.get("extraction_method", "unknown"),
                    "table_confidence": block.get("table_confidence", 0.0),
                    "image_size": block.get("image_size"),
                    "image_index": block.get("image_index")
                })
            elif block["type"] == "image_chart":
                metadata.update({
                    "source_type": "chart",
                    "block_type": "chart",
                    "chart_type": block.get("chart_type", "unknown"),
                    "chart_confidence": block.get("chart_confidence", 0.0),
                    "structured_data": block.get("structured_data", []),
                    "image_size": block.get("image_size"),
                    "image_index": block.get("image_index")
                })
            elif block["type"] == "text":
                metadata.update({
                    "style_name": block.get("style_name", "Normal")
                })
            elif block["type"] == "image_reference":
                metadata.update({
                    "source_type": "image",
                    "image_size": block.get("image_size"),
                    "image_index": block.get("image_index")
                })

            processed_blocks.append((block["text"], metadata))

        # Process actual images if image processor is available
        if self.image_processor:
            image_blocks = self._extract_images_from_docx(container, blob_name)
            for i, (img_text, img_metadata) in enumerate(image_blocks):
                img_metadata["doc_element_order"] = len(processed_blocks) + 1 + i
                if "page_numbers" not in img_metadata:
                    img_metadata["page_numbers"] = [1]
                processed_blocks.append((img_text, img_metadata))

        if not processed_blocks:
            logger.warning(f"No content blocks extracted from {blob_name}")

        logger.info(f"Successfully processed {blob_name} (enhanced): {len(processed_blocks)} blocks extracted")
        return processed_blocks

    def _determine_source_type(self, block_type: str) -> str:
        """Determine source type from block type."""
        type_mapping = {
            "table": "table", "image_table": "table", "image_chart": "chart",
            "chart": "chart", "image_reference": "image", "text": "text"
        }
        return type_mapping.get(block_type, "text")

    def _process_file_legacy(self, container: str, blob_name: str) -> List[Tuple[str, Dict[str, Any]]]:
        """
        Process file using legacy logic for backward compatibility.

        Args:
            container: Azure Blob Storage container name
            blob_name: Blob name for the DOCX file

        Returns:
            List of (text_block, metadata) tuples
        """
        processed_blocks: List[Tuple[str, Dict[str, Any]]] = []
        element_order = 0

        # Load document from Azure Blob Storage
        document = blob_storage_service.get_docx_document(container, blob_name)

        # Process document elements in order
        for element in document.element.body:
            element_order += 1
            page_numbers = self._get_page_numbers_from_element(element, document)

            # Process paragraphs
            if hasattr(element, 'tag') and 'p' in element.tag:
                try:
                    paragraph = docx.text.paragraph.Paragraph(element, document)
                    result = self._process_paragraph(paragraph, element_order, page_numbers)
                    if result:
                        processed_blocks.append(result)
                except:
                    pass

            # Process tables
            elif hasattr(element, 'tag') and 'tbl' in element.tag:
                try:
                    table = docx.table.Table(element, document)
                    result = self._process_table(table, element_order, page_numbers)
                    if result:
                        processed_blocks.append(result)
                except:
                    pass

        # Extract and process images
        image_blocks = self._extract_images_from_docx(container, blob_name)
        for i, (img_text, img_metadata) in enumerate(image_blocks):
            img_metadata["doc_element_order"] = element_order + 1 + i
            if "page_numbers" not in img_metadata:
                img_metadata["page_numbers"] = [1]
            processed_blocks.append((img_text, img_metadata))

        if not processed_blocks:
            logger.warning(f"No content blocks extracted from {blob_name}")

        logger.info(f"Successfully processed {blob_name} (legacy): {len(processed_blocks)} blocks extracted")
        return processed_blocks

    def save_metadata_and_blocks(self, container: str, blob_name: str, output_container: str = None, blocks_input: Any = None) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
        """
        Process file and save metadata and blocks to Azure Blob Storage.

        Args:
            container: Source Azure Blob Storage container name
            blob_name: Source blob name for the DOCX file
            output_container: Optional separate container for output files
            blocks_input: Optional pre-parsed blocks (can handle various formats)

        Returns:
            Tuple of (blocks, metadata)
        """
        # Use same container for output if not specified
        if output_container is None:
            output_container = container

        # Parse using enhanced logic if blocks not provided
        if blocks_input is None:
            blocks, image_info = self._parse_docx_to_blocks_enhanced(container, blob_name)
        else:
            # Use provided blocks and normalize them
            blocks = self._normalize_blocks_input(blocks_input)
            image_info = []

        # Create LlamaIndex nodes if available
        nodes = self.create_llamaindex_nodes(blocks) if self.use_llamaindex else []

        # Generate comprehensive metadata
        metadata = self.generate_metadata_with_llamaindex(container, blob_name, blocks, nodes)

        # Generate output file names
        file_stem = Path(blob_name).stem
        metadata_blob_name = f"{file_stem}_metadata.json"
        blocks_blob_name = f"{file_stem}_parsed_blocks.json"

        try:
            # Save metadata JSON to Azure Blob Storage
            metadata_json = json.dumps(metadata, ensure_ascii=False, indent=2)
            blob_storage_service.upload_text(output_container, metadata_blob_name, metadata_json)
            logger.info(f"Metadata saved to: {output_container}/{metadata_blob_name}")

            # Save parsed blocks JSON to Azure Blob Storage
            blocks_json = json.dumps(blocks, ensure_ascii=False, indent=2)
            blob_storage_service.upload_text(output_container, blocks_blob_name, blocks_json)
            logger.info(f"Parsed blocks saved to: {output_container}/{blocks_blob_name}")

        except Exception as e:
            logger.error(f"Error saving metadata and blocks to Azure Blob Storage: {e}")
            # Continue execution even if saving fails
            pass

        return blocks, metadata

    def process_docx(self, container: str, blob_name: str) -> List[Tuple[str, Dict[str, Any]]]:
        """
        Legacy method for backward compatibility.

        Args:
            container: Azure Blob Storage container name
            blob_name: Blob name for the DOCX file

        Returns:
            List of (text_block, metadata) tuples
        """
        return self.process_file(container, blob_name)

    def get_processing_stats(self, processed_blocks: List[Tuple[str, Dict[str, Any]]]) -> Dict[str, Any]:
        """
        Get comprehensive statistics about the processing results.

        Args:
            processed_blocks: List of processed blocks

        Returns:
            Dictionary with processing statistics
        """
        stats = {
            "total_blocks": len(processed_blocks),
            "block_types": {},
            "source_types": {},
            "extraction_methods": {},
            "average_text_length": 0,
            "tables_detected": 0,
            "charts_detected": 0,
            "images_processed": 0
        }

        total_text_length = 0
        
        for text, metadata in processed_blocks:
            # Count block types
            block_type = metadata.get("block_type", "unknown")
            stats["block_types"][block_type] = stats["block_types"].get(block_type, 0) + 1
            
            # Count source types
            source_type = metadata.get("source_type", "unknown")
            stats["source_types"][source_type] = stats["source_types"].get(source_type, 0) + 1
            
            # Count extraction methods
            extraction_method = metadata.get("extraction_method", "standard")
            stats["extraction_methods"][extraction_method] = stats["extraction_methods"].get(extraction_method, 0) + 1
            
            # Calculate text statistics
            total_text_length += len(text)
            
            # Count specific content types
            if block_type == "table" or source_type == "table":
                stats["tables_detected"] += 1
            if block_type == "chart" or source_type == "chart":
                stats["charts_detected"] += 1
            if source_type == "image":
                stats["images_processed"] += 1
        
        stats["average_text_length"] = total_text_length / len(processed_blocks) if processed_blocks else 0
        return stats

    def validate_processing_output(self, processed_blocks: List[Tuple[str, Dict[str, Any]]]) -> Dict[str, Any]:
        """
        Validate the processing output for quality and completeness.

        Args:
            processed_blocks: List of processed blocks

        Returns:
            Validation results dictionary
        """
        validation_results = {
            "is_valid": True,
            "errors": [],
            "warnings": [],
            "total_blocks": len(processed_blocks)
        }

        for i, (text, metadata) in enumerate(processed_blocks):
            # Validate text content
            if not isinstance(text, str):
                validation_results["errors"].append(f"Block {i}: text is not a string")
                validation_results["is_valid"] = False
            
            # Validate metadata structure
            if not isinstance(metadata, dict):
                validation_results["errors"].append(f"Block {i}: metadata is not a dictionary")
                validation_results["is_valid"] = False
                continue
            
            # Check required metadata fields
            required_fields = ["block_type", "source_type", "file_type", "page_numbers"]
            for field in required_fields:
                if field not in metadata:
                    validation_results["warnings"].append(f"Block {i}: missing {field} in metadata")
            
            # Check for empty content
            if not text.strip():
                validation_results["warnings"].append(f"Block {i}: empty text content")
            
            # Validate table-specific metadata
            if metadata.get("block_type") == "table":
                if "num_rows" not in metadata and "table_confidence" not in metadata:
                    validation_results["warnings"].append(f"Block {i}: table missing row count or confidence")
            
            # Validate page numbers
            page_nums = metadata.get("page_numbers", [])
            if not isinstance(page_nums, list) or not page_nums:
                validation_results["warnings"].append(f"Block {i}: invalid or missing page_numbers")

        return validation_results